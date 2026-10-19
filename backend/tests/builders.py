"""Small test documents built on the fly, so no binary file is stored in Git."""

from pathlib import Path

import docx
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

CONTRACT_TEXT = "The termination notice period is thirty (30) days."


def make_scan(path: Path, text: str) -> Path:
    """Save a PNG that looks like a clean scan of printed text."""
    image = Image.new("RGB", (1600, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 150), text, fill="black", font=ImageFont.load_default(size=56))
    image.save(path)
    return path


def make_pdf(path: Path, pages: list[str | Path]) -> Path:
    """Build a PDF: a str becomes a text page, a Path becomes an image-only page."""
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    for content in pages:
        pdf.add_page()
        if isinstance(content, Path):
            pdf.image(str(content), x=10, y=10, w=190)
        else:
            pdf.multi_cell(0, 8, content)
    pdf.output(str(path))
    return path


def make_docx(path: Path) -> Path:
    """Build a DOCX with an introduction, two headings and a table."""
    document = docx.Document()
    document.add_paragraph("Service Agreement between Acme Corp and Beta LLC.")
    document.add_heading("1. Term", level=1)
    document.add_paragraph("This agreement lasts twelve months.")
    document.add_heading("2. Fees", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Monthly fee"
    table.cell(1, 0).text = "Support"
    table.cell(1, 1).text = "USD 1,200"
    document.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    """Build a PPTX with a title slide, a slide with a table, and an empty slide."""
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "Quarterly Business Review"
    title_slide.placeholders[1].text = "Q3 2026"

    blank_layout = presentation.slide_layouts[6]
    table_slide = presentation.slides.add_slide(blank_layout)
    heading = table_slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
    heading.text_frame.text = "Revenue by region"
    rows = [("Region", "Revenue"), ("EMEA", "USD 4.2M")]
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value

    presentation.slides.add_slide(blank_layout)  # empty slide
    presentation.save(str(path))
    return path