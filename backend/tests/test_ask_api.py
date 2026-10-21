"""Tests for POST /api/ask (fake embeddings, fake LLM, real Chroma)."""

import io
from collections.abc import Callable, Iterator

import docx
import pytest
from fastapi.testclient import TestClient

from app.api.dependencies import get_indexer
from app.core.config import get_settings
from app.main import app
from app.rag.answering import REFUSAL_MESSAGE
from app.retrieval.indexing import DocumentIndexer

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
QUESTION = "How can a party terminate the agreement?"


def make_docx() -> bytes:
    """Return a Word document with two heading-based sections."""
    document = docx.Document()
    document.add_heading("Start date", level=1)
    document.add_paragraph("This service agreement starts on 1 January 2026 in Paris.")
    document.add_heading("Termination", level=1)
    document.add_paragraph("Either party may terminate the agreement with 30 days notice.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload_contract(client: TestClient) -> None:
    upload = client.post(
        "/api/documents", files={"file": ("contract.docx", make_docx(), DOCX_MIME)}
    )
    assert upload.json()["status"] == "ready"


@pytest.fixture
def set_min_score(client: TestClient) -> Iterator[Callable[[float], None]]:
    """Change the relevance threshold for one test.

    The fake embedding model does not give the same scores as nomic-embed-text,
    so tests choose the threshold instead of relying on the default value.
    """

    def apply(value: float) -> None:
        settings = get_settings().model_copy(update={"retrieval_min_score": value})
        app.dependency_overrides[get_settings] = lambda: settings

    yield apply
    app.dependency_overrides.pop(get_settings, None)


def test_ask_returns_an_answer_with_checked_sources(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.0)
    upload_contract(client)
    fake_llm.replies.append("Either party may terminate with 30 days notice [S1].")

    response = client.post("/api/ask", json={"question": QUESTION, "top_k": 1})

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "answered"
    assert body["refusal_reason"] is None
    assert len(body["citations"]) == 1
    citation = body["citations"][0]
    assert citation["label"] == "S1"
    assert citation["filename"] == "contract.docx"
    assert "terminate" in citation["text"]
    assert citation["display"] == f"[contract.docx, p. {citation['page_number']}]"
    expected = f"Either party may terminate with 30 days notice {citation['display']}."
    assert body["answer"] == expected
    assert body["warnings"] == []
    assert body["invalid_citations"] == []
    assert isinstance(body["best_score"], float)
    assert [(run["name"], run["kind"], run["status"]) for run in body["agent_trace"]] == [
        ("retrieval", "tool", "ok"),
        ("writer", "agent", "ok"),
    ]
    assert all(run["duration_ms"] >= 0 for run in body["agent_trace"])
    # The model received the question and the passage, but no file name.
    prompt = fake_llm.calls[0].messages[1].content
    assert QUESTION in prompt
    assert "contract.docx" not in prompt


def test_ask_refuses_without_calling_the_model_below_the_threshold(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.999)
    upload_contract(client)

    response = client.post("/api/ask", json={"question": "What is the CEO's salary?"})

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "refused"
    assert body["refusal_reason"] == "no_relevant_passages"
    assert body["answer"] == REFUSAL_MESSAGE
    assert body["citations"] == []
    assert body["best_score"] is not None
    assert fake_llm.calls == []
    assert [(run["name"], run["status"]) for run in body["agent_trace"]] == [
        ("retrieval", "ok"),
        ("writer", "skipped"),
    ]


def test_ask_with_no_document_refuses_without_calling_the_model(
    client: TestClient, fake_llm
) -> None:
    response = client.post("/api/ask", json={"question": QUESTION})

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "refused"
    assert body["refusal_reason"] == "no_relevant_passages"
    assert body["best_score"] is None
    assert fake_llm.calls == []


def test_ask_reports_a_refusal_by_the_model(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.0)
    upload_contract(client)
    fake_llm.replies.append("NOT_FOUND")

    body = client.post("/api/ask", json={"question": "What is the CEO's salary?"}).json()

    assert body["status"] == "refused"
    assert body["refusal_reason"] == "not_found_by_model"
    assert body["answer"] == REFUSAL_MESSAGE


def test_ask_returns_503_when_the_model_is_down(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.0)
    upload_contract(client)
    fake_llm.unavailable = True

    response = client.post("/api/ask", json={"question": QUESTION})

    assert response.status_code == 503
    assert response.json()["detail"].startswith("Answer generation failed")


def test_ask_returns_503_when_the_embedding_service_is_down(
    client: TestClient, make_indexer: Callable[..., DocumentIndexer]
) -> None:
    app.dependency_overrides[get_indexer] = lambda: make_indexer(unavailable=True, name="down")

    response = client.post("/api/ask", json={"question": QUESTION})

    assert response.status_code == 503
    assert "Search failed" in response.json()["detail"]


@pytest.mark.parametrize(
    "payload",
    [
        {},
        {"question": "   "},
        {"question": "x" * 1001},
        {"question": "ok", "top_k": 0},
        {"question": "ok", "top_k": 21},
    ],
)
def test_ask_refuses_invalid_requests(client: TestClient, payload: dict) -> None:
    assert client.post("/api/ask", json=payload).status_code == 422