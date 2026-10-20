"""Tests for the retriever and POST /api/search (fake embeddings, real Chroma)."""

import io
from collections.abc import Callable

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.api.dependencies import get_indexer
from app.db.repositories import DocumentRepository
from app.ingestion.models import PageText
from app.main import app
from app.retrieval.indexing import DocumentIndexer
from app.retrieval.retriever import RetrievalError, Retriever
from app.retrieval.vector_store import VectorRecord

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

CONTRACT_PAGES = [
    PageText(1, "This service agreement starts on 1 January 2026 in Paris."),
    PageText(2, "Either party may terminate the agreement with 30 days written notice."),
    PageText(3, "Payment is due within 45 days of the invoice date."),
]


def retriever_for(indexer: DocumentIndexer) -> Retriever:
    return Retriever(indexer.embedder, indexer.vector_store)


def make_docx() -> bytes:
    """Return a Word document with two heading-based sections."""
    document = docx.Document()
    document.add_heading("Start date", level=1)
    document.add_paragraph("This service agreement starts on 1 January 2026 in Paris.")
    document.add_heading("Termination", level=1)
    document.add_paragraph("Either party may terminate the agreement with 30 days notice.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_search_returns_the_right_page_with_its_source(
    db_session: Session, indexer: DocumentIndexer
) -> None:
    repository = DocumentRepository(db_session)
    document = repository.create("contract.pdf", "application/pdf")
    indexer.index(repository, document, CONTRACT_PAGES)

    results = retriever_for(indexer).search(
        repository, "How can a party terminate the agreement?", top_k=2
    )

    assert len(results) == 2
    best = results[0]
    assert (best.filename, best.page_number, best.position) == ("contract.pdf", 2, 0)
    assert best.text == CONTRACT_PAGES[1].text
    assert best.document_id == document.id
    assert best.score > results[1].score


def test_search_can_be_limited_to_some_documents(
    db_session: Session, indexer: DocumentIndexer
) -> None:
    repository = DocumentRepository(db_session)
    contract = repository.create("contract.pdf", "application/pdf")
    invoice = repository.create("invoice.pdf", "application/pdf")
    indexer.index(repository, contract, CONTRACT_PAGES)
    indexer.index(repository, invoice, [PageText(1, "Invoice total: 1,200 EUR, due in 45 days.")])

    results = retriever_for(indexer).search(
        repository, "When is payment due?", top_k=5, document_ids=[invoice.id]
    )

    assert [result.filename for result in results] == ["invoice.pdf"]


def test_vectors_without_a_chunk_in_sqlite_are_skipped(
    db_session: Session, indexer: DocumentIndexer
) -> None:
    repository = DocumentRepository(db_session)
    document = repository.create("contract.pdf", "application/pdf")
    indexer.index(repository, document, CONTRACT_PAGES)
    question = "terminate the agreement notice"
    indexer.vector_store.upsert(
        [
            VectorRecord(
                id="orphan-id",
                embedding=indexer.embedder.embed_query(question),
                document_id=document.id,
                page_number=9,
                position=0,
            )
        ]
    )

    results = retriever_for(indexer).search(repository, question, top_k=5)

    assert "orphan-id" not in {result.chunk_id for result in results}
    assert len(results) == 3


def test_search_errors_are_reported_as_retrieval_errors(
    db_session: Session, make_indexer: Callable[..., DocumentIndexer]
) -> None:
    broken = make_indexer(unavailable=True)
    with pytest.raises(RetrievalError, match="Search failed"):
        retriever_for(broken).search(DocumentRepository(db_session), "Any question?", top_k=5)


def test_search_endpoint_finds_the_uploaded_passage(client: TestClient) -> None:
    upload = client.post(
        "/api/documents", files={"file": ("contract.docx", make_docx(), DOCX_MIME)}
    )
    assert upload.json()["status"] == "ready"

    response = client.post(
        "/api/search", json={"query": "How can a party terminate the agreement?", "top_k": 1}
    )

    assert response.status_code == 200
    body = response.json()
    assert body["query"] == "How can a party terminate the agreement?"
    assert len(body["results"]) == 1
    best = body["results"][0]
    assert best["filename"] == "contract.docx"
    assert "terminate" in best["text"]
    assert best["document_id"] == upload.json()["id"]


def test_search_endpoint_uses_the_default_top_k(client: TestClient) -> None:
    client.post("/api/documents", files={"file": ("contract.docx", make_docx(), DOCX_MIME)})

    response = client.post("/api/search", json={"query": "agreement"})

    assert response.status_code == 200
    assert 1 <= len(response.json()["results"]) <= 5


def test_search_endpoint_with_no_document_returns_no_result(client: TestClient) -> None:
    response = client.post("/api/search", json={"query": "termination notice period"})
    assert response.status_code == 200
    assert response.json()["results"] == []


@pytest.mark.parametrize(
    "payload",
    [
        {"query": "   "},
        {"query": ""},
        {"query": "x" * 1001},
        {"query": "ok", "top_k": 0},
        {"query": "ok", "top_k": 21},
    ],
)
def test_search_endpoint_refuses_invalid_requests(client: TestClient, payload: dict) -> None:
    assert client.post("/api/search", json=payload).status_code == 422


def test_search_endpoint_returns_503_when_the_embedding_service_is_down(
    client: TestClient, make_indexer: Callable[..., DocumentIndexer]
) -> None:
    app.dependency_overrides[get_indexer] = lambda: make_indexer(unavailable=True, name="down")

    response = client.post("/api/search", json={"query": "termination notice period"})

    assert response.status_code == 503
    assert "Search failed" in response.json()["detail"]