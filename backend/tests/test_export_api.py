"""Tests for GET /api/sessions/{id}/messages/{id}/export (no LLM call, real database)."""

import io
from collections.abc import Iterator
from datetime import UTC, datetime
from typing import Any

import docx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.api.export import get_export_time
from app.db.models import MessageMode
from app.db.repositories import AssistantReply, SessionRepository
from app.main import app

NOW = datetime(2026, 10, 23, 19, 20, tzinfo=UTC)
QUESTION = "How much was invoiced in Q2 2026?"
REVIEW_FILE = "quarterly_review_HDS_Q2_2026.pptx"
ANSWER = f"The total invoiced in Q2 2026 was USD 45,370.00 [{REVIEW_FILE}, p. 5]."
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture
def fixed_time(client: TestClient) -> Iterator[None]:
    app.dependency_overrides[get_export_time] = lambda: NOW
    yield
    app.dependency_overrides.pop(get_export_time, None)


def citation(page: int, text: str, filename: str = REVIEW_FILE) -> dict[str, Any]:
    return {
        "label": "S1",
        "display": f"[{filename}, p. {page}]",
        "chunk_id": f"chunk-{page}",
        "document_id": "doc-1",
        "filename": filename,
        "page_number": page,
        "position": 0,
        "text": text,
        "score": 0.8,
    }


ANSWERED = AssistantReply(
    content=ANSWER,
    status="answered",
    sources=[citation(5, "Total invoiced in the quarter: USD 45,370.00.")],
    best_score=0.8,
    review={"verdict": "ok", "issues": [], "model_checked": True},
)
REFUSED = AssistantReply(
    content="The documents do not contain this information.",
    status="refused",
    refusal_reason="no_relevant_passages",
    best_score=0.5,
)
INVOICE_FILE = "invoice_HDS-INV-2026-0042.pdf"
EXTRACTED = AssistantReply(
    content="Extracted 1 of 1 fields.",
    status="answered",
    review={"verdict": "ok", "issues": [], "model_checked": False},
    extraction={
        "document_id": "doc-2",
        "filename": INVOICE_FILE,
        "document_type": "invoice",
        "fields": [
            {
                "name": "invoice_number",
                "label": "Invoice number",
                "value": "HDS-INV-2026-0042",
                "page": 1,
                "source": citation(1, "Invoice No.: HDS-INV-2026-0042", INVOICE_FILE),
            }
        ],
    },
)


def save_exchange(
    question: str, reply: AssistantReply, mode: MessageMode = MessageMode.ASK
) -> tuple[str, str, str]:
    """Save a question and its reply directly; return (session, question, answer) ids."""
    with app.state.session_factory() as db:
        repository = SessionRepository(db)
        chat_session = repository.create()
        asked, answered = repository.add_exchange(
            chat_session, question=question, document_ids=None, reply=reply, mode=mode
        )
        return chat_session.id, asked.id, answered.id


def export(
    client: TestClient,
    session_id: str,
    message_id: str,
    file_format: str,
    language: str | None = None,
):
    params = {"format": file_format} | ({"lang": language} if language else {})
    return client.get(f"/api/sessions/{session_id}/messages/{message_id}/export", params=params)


def test_an_answer_exports_as_word(client: TestClient, fixed_time: None, fake_llm) -> None:
    session_id, _, answer_id = save_exchange(QUESTION, ANSWERED)

    response = export(client, session_id, answer_id, "docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="ATLAS_answer_2026-10-23.docx"'
    )
    document = docx.Document(io.BytesIO(response.content))
    lines = [paragraph.text for paragraph in document.paragraphs]
    assert QUESTION in lines
    assert ANSWER in lines
    assert f"[{REVIEW_FILE}, p. 5]" in lines
    assert "Total invoiced in the quarter: USD 45,370.00." in lines
    assert "23 October 2026, 19:20 UTC" in document.sections[0].footer.paragraphs[0].text
    # The file is built from the database: the model is never called.
    assert fake_llm.calls == []


def test_an_answer_exports_as_powerpoint(client: TestClient, fixed_time: None) -> None:
    session_id, _, answer_id = save_exchange(QUESTION, ANSWERED)

    response = export(client, session_id, answer_id, "pptx")

    assert response.status_code == 200
    assert response.headers["content-type"] == PPTX_TYPE
    assert 'filename="ATLAS_answer_2026-10-23.pptx"' in response.headers["content-disposition"]
    presentation = Presentation(io.BytesIO(response.content))
    assert len(presentation.slides) == 3
    assert presentation.core_properties.subject == QUESTION


def test_an_extraction_exports_with_its_table(client: TestClient, fixed_time: None) -> None:
    session_id, _, answer_id = save_exchange(
        f"Extract the key fields of {INVOICE_FILE}", EXTRACTED, MessageMode.EXTRACT
    )

    response = export(client, session_id, answer_id, "docx")

    assert response.status_code == 200
    assert 'filename="ATLAS_extraction_2026-10-23.docx"' in response.headers["content-disposition"]
    table = docx.Document(io.BytesIO(response.content)).tables[0]
    assert [cell.text for cell in table.rows[1].cells] == [
        "Invoice number",
        "HDS-INV-2026-0042",
        "p. 1",
    ]


def test_the_question_of_the_answer_is_the_one_exported(client: TestClient) -> None:
    """In a session of two exchanges, each answer comes with its own question."""
    session_id, _, first_id = save_exchange(QUESTION, ANSWERED)
    with app.state.session_factory() as db:
        repository = SessionRepository(db)
        chat_session = repository.get(session_id)
        assert chat_session is not None
        _, second = repository.add_exchange(
            chat_session, question="And in Q1?", document_ids=None, reply=ANSWERED
        )
        second_id = second.id

    first = docx.Document(io.BytesIO(export(client, session_id, first_id, "docx").content))
    second_file = docx.Document(io.BytesIO(export(client, session_id, second_id, "docx").content))

    assert QUESTION in [p.text for p in first.paragraphs]
    assert "And in Q1?" in [p.text for p in second_file.paragraphs]
    assert QUESTION not in [p.text for p in second_file.paragraphs]


def test_a_question_or_a_refusal_gives_409(client: TestClient) -> None:
    session_id, question_id, refused_id = save_exchange("What is the CEO's salary?", REFUSED)

    for message_id in (question_id, refused_id):
        response = export(client, session_id, message_id, "docx")
        assert response.status_code == 409
        assert response.json()["detail"] == "Only an answer or an extraction can be exported."


def test_unknown_session_or_message_gives_404(client: TestClient) -> None:
    session_id, _, answer_id = save_exchange(QUESTION, ANSWERED)
    other_session, _, _ = save_exchange("Another question?", ANSWERED)

    unknown_session = export(client, "does-not-exist", answer_id, "docx")
    unknown_message = export(client, session_id, "does-not-exist", "docx")
    # A real message, asked through the wrong session.
    wrong_session = export(client, other_session, answer_id, "docx")

    assert unknown_session.status_code == 404
    assert unknown_session.json()["detail"] == "Session not found."
    assert unknown_message.status_code == 404
    assert unknown_message.json()["detail"] == "Message not found."
    assert wrong_session.status_code == 404


def test_a_missing_or_unknown_format_gives_422(client: TestClient) -> None:
    session_id, _, answer_id = save_exchange(QUESTION, ANSWERED)
    url = f"/api/sessions/{session_id}/messages/{answer_id}/export"

    assert client.get(url).status_code == 422
    assert client.get(url, params={"format": "pdf"}).status_code == 422


def test_an_answer_exports_in_french(client: TestClient, fixed_time: None) -> None:
    session_id, _, answer_id = save_exchange(QUESTION, ANSWERED)

    response = export(client, session_id, answer_id, "docx", "fr")

    assert response.status_code == 200
    # The file name stays the same in both languages (no accents in the header).
    assert 'filename="ATLAS_answer_2026-10-23.docx"' in response.headers["content-disposition"]
    document = docx.Document(io.BytesIO(response.content))
    lines = [paragraph.text for paragraph in document.paragraphs]
    assert "Réponse tirée de vos documents" in lines
    # What was asked and answered is kept as written.
    assert QUESTION in lines
    assert ANSWER in lines
    assert "23 octobre 2026, 19:20 UTC" in document.sections[0].footer.paragraphs[0].text


def test_an_extraction_exports_in_french_as_powerpoint(
    client: TestClient, fixed_time: None
) -> None:
    session_id, _, answer_id = save_exchange(
        f"Extract the key fields of {INVOICE_FILE}", EXTRACTED, MessageMode.EXTRACT
    )

    response = export(client, session_id, answer_id, "pptx", "fr")

    assert response.status_code == 200
    presentation = Presentation(io.BytesIO(response.content))
    texts = [
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    cells = [
        cell.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    ]
    assert f"Champs clés de {INVOICE_FILE} (facture)" in texts
    assert "Numéro de facture" in cells
    assert "HDS-INV-2026-0042" in cells


def test_english_is_the_default_and_an_unknown_language_gives_422(client: TestClient) -> None:
    session_id, _, answer_id = save_exchange(QUESTION, ANSWERED)

    default = docx.Document(io.BytesIO(export(client, session_id, answer_id, "docx").content))
    english = docx.Document(io.BytesIO(export(client, session_id, answer_id, "docx", "en").content))

    assert "Answer from your documents" in [p.text for p in default.paragraphs]
    assert [p.text for p in default.paragraphs] == [p.text for p in english.paragraphs]
    assert export(client, session_id, answer_id, "docx", "de").status_code == 422