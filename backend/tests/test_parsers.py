"""Tests for the PDF, DOCX, PPTX and image parsers.

Every test builds its own small document, so no binary file is stored in Git.
Tests marked "needs_tesseract" are skipped where Tesseract is not installed.
"""

from pathlib import Path

import docx
import pytest
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from app.ingestion.models import FileType, PageText, ParseError, ParseOptions
from app.ingestion.ocr import is_tesseract_available
from app.ingestion.parsing import parse_document

needs_tesseract = pytest.mark.skipif(
    not is_tesseract_available(), reason="Tesseract is not installed"
)

OPTIONS = ParseOptions()
CONTRACT_TEXT = "The termination notice period is thirty (30) days."


# --- Helpers that build test documents ---------------------------------------


def make_scan(path: Path, text: str) -> Path:
    """Save a PNG that looks like a clean scan of printed text."""
    image = Image.new("RGB", (1600, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 150), text, fill="black", font=ImageFont.load_default(size=56))
    image.save(path)
    return path


def make_pdf(path: Path, pages: list[str | Path]) -> Path:
    """Build a PDF: a str becomes a text page, a Path becomes an image-only page."""
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    for content in pages:
        pdf.add_page()
        if isinstance(content, Path):
            pdf.image(str(content), x=10, y=10, w=190)
        else:
            pdf.multi_cell(0, 8, content)
    pdf.output(str(path))
    return path


def make_docx(path: Path) -> Path:
    """Build a DOCX with an introduction, two headings and a table."""
    document = docx.Document()
    document.add_paragraph("Service Agreement between Acme Corp and Beta LLC.")
    document.add_heading("1. Term", level=1)
    document.add_paragraph("This agreement lasts twelve months.")
    document.add_heading("2. Fees", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Monthly fee"
    table.cell(1, 0).text = "Support"
    table.cell(1, 1).text = "USD 1,200"
    document.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    """Build a PPTX with a title slide, a slide with a table, and an empty slide."""
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "Quarterly Business Review"
    title_slide.placeholders[1].text = "Q3 2026"

    blank_layout = presentation.slide_layouts[6]
    table_slide = presentation.slides.add_slide(blank_layout)
    heading = table_slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
    heading.text_frame.text = "Revenue by region"
    rows = [("Region", "Revenue"), ("EMEA", "USD 4.2M")]
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value

    presentation.slides.add_slide(blank_layout)  # empty slide
    presentation.save(str(path))
    return path


# --- PDF ---------------------------------------------------------------------


def test_pdf_text_pages_are_read_without_ocr(tmp_path: Path) -> None:
    """A PDF with a text layer is read directly, page by page."""
    pdf = make_pdf(tmp_path / "contract.pdf", [CONTRACT_TEXT, "Payment is due within 15 days."])

    pages = parse_document(pdf, FileType.PDF, OPTIONS)

    assert [page.page_number for page in pages] == [1, 2]
    assert pages[0].text == CONTRACT_TEXT
    assert "15 days" in pages[1].text
    assert not any(page.ocr_used for page in pages)


@needs_tesseract
def test_scanned_pdf_page_goes_through_ocr(tmp_path: Path) -> None:
    """A page that is only an image is read with OCR; other pages are not."""
    scan = make_scan(tmp_path / "scan.png", "INVOICE NUMBER 2026-0042")
    pdf = make_pdf(tmp_path / "mixed.pdf", [CONTRACT_TEXT, scan])

    pages = parse_document(pdf, FileType.PDF, OPTIONS)

    assert pages[0].ocr_used is False
    assert pages[1].ocr_used is True
    assert "2026-0042" in pages[1].text


def test_pdf_over_the_page_limit_is_refused(tmp_path: Path) -> None:
    """Long documents are refused before any text is extracted."""
    pdf = make_pdf(tmp_path / "long.pdf", ["Page text that is long enough."] * 3)

    with pytest.raises(ParseError, match="3 pages; the limit is 2"):
        parse_document(pdf, FileType.PDF, ParseOptions(max_pages=2))


def test_damaged_pdf_raises_parse_error(tmp_path: Path) -> None:
    """A file that is not a real PDF gives a clear error."""
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.7 this is not really a PDF")

    with pytest.raises(ParseError, match="Cannot open the PDF"):
        parse_document(broken, FileType.PDF, OPTIONS)


def test_password_protected_pdf_raises_parse_error(tmp_path: Path) -> None:
    """Protected PDFs are refused with a message that mentions the password."""
    pdf = FPDF()
    pdf.set_encryption(owner_password="owner", user_password="secret")
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(0, 10, CONTRACT_TEXT)
    protected = tmp_path / "protected.pdf"
    pdf.output(str(protected))

    with pytest.raises(ParseError, match="password"):
        parse_document(protected, FileType.PDF, OPTIONS)


# --- DOCX --------------------------------------------------------------------


def test_docx_is_split_into_heading_sections(tmp_path: Path) -> None:
    """Each heading starts a new numbered section; tables keep their rows."""
    pages = parse_document(make_docx(tmp_path / "agreement.docx"), FileType.DOCX, OPTIONS)

    assert pages == [
        PageText(1, "Service Agreement between Acme Corp and Beta LLC."),
        PageText(2, "1. Term\nThis agreement lasts twelve months."),
        PageText(3, "2. Fees\nService | Monthly fee\nSupport | USD 1,200"),
    ]


def test_docx_without_headings_is_one_section(tmp_path: Path) -> None:
    """A document without headings becomes a single section."""
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    path = tmp_path / "plain.docx"
    document.save(str(path))

    pages = parse_document(path, FileType.DOCX, OPTIONS)

    assert pages == [PageText(1, "First paragraph.\nSecond paragraph.")]


def test_damaged_docx_raises_parse_error(tmp_path: Path) -> None:
    """A file that is not a real DOCX gives a clear error."""
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip file")

    with pytest.raises(ParseError, match="Cannot open the DOCX"):
        parse_document(broken, FileType.DOCX, OPTIONS)


# --- PPTX --------------------------------------------------------------------


def test_pptx_slides_become_pages(tmp_path: Path) -> None:
    """One slide is one page; tables are read; empty slides keep their number."""
    pages = parse_document(make_pptx(tmp_path / "review.pptx"), FileType.PPTX, OPTIONS)

    assert [page.page_number for page in pages] == [1, 2, 3]
    assert pages[0].text == "Quarterly Business Review\nQ3 2026"
    assert pages[1].text == "Revenue by region\nRegion | Revenue\nEMEA | USD 4.2M"
    assert pages[2].text == ""


def test_pptx_group_shapes_are_read(tmp_path: Path) -> None:
    """Text inside grouped shapes is not lost."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Text inside a group"
    path = tmp_path / "group.pptx"
    presentation.save(str(path))

    pages = parse_document(path, FileType.PPTX, OPTIONS)

    assert pages == [PageText(1, "Text inside a group")]


def test_damaged_pptx_raises_parse_error(tmp_path: Path) -> None:
    """A file that is not a real PPTX gives a clear error."""
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a zip file")

    with pytest.raises(ParseError, match="Cannot open the PPTX"):
        parse_document(broken, FileType.PPTX, OPTIONS)


# --- Images ------------------------------------------------------------------


@needs_tesseract
def test_image_is_read_with_ocr_as_page_one(tmp_path: Path) -> None:
    """An image is one page, always read with OCR."""
    scan = make_scan(tmp_path / "invoice.png", "TOTAL DUE USD 3,450")

    pages = parse_document(scan, FileType.IMAGE, OPTIONS)

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].ocr_used is True
    assert "3,450" in pages[0].text


def test_damaged_image_raises_parse_error(tmp_path: Path) -> None:
    """A file that is not a real image gives a clear error."""
    broken = tmp_path / "broken.png"
    broken.write_bytes(b"not an image")

    with pytest.raises(ParseError, match="Cannot open the image"):
        parse_document(broken, FileType.IMAGE, OPTIONS)


# --- Common rules ------------------------------------------------------------


@needs_tesseract
def test_document_without_any_text_is_refused(tmp_path: Path) -> None:
    """A blank document fails clearly instead of being stored as empty."""
    blank = tmp_path / "blank.png"
    Image.new("RGB", (800, 600), "white").save(blank)

    with pytest.raises(ParseError, match="No text could be extracted"):
        parse_document(blank, FileType.IMAGE, OPTIONS)