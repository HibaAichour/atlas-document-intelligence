"""Tests for the indexer and for indexing during upload (fake embeddings, real Chroma)."""

import io
from collections.abc import Callable

import docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.api.dependencies import get_indexer
from app.db.repositories import DocumentRepository
from app.ingestion.models import PageText
from app.main import app
from app.retrieval.indexing import DocumentIndexer, IndexingError

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

CONTRACT_PAGES = [
    PageText(1, "This service agreement starts on 1 January 2026."),
    PageText(2, "Either party may terminate the agreement with 30 days written notice."),
    PageText(3, "   "),  # a blank page gives no chunk
]


def make_docx() -> bytes:
    """Return a small Word document with one heading and one paragraph."""
    document = docx.Document()
    document.add_heading("Termination", level=1)
    document.add_paragraph("Either party may terminate the agreement with 30 days notice.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_index_stores_chunks_and_vectors_under_the_same_ids(
    db_session: Session, indexer: DocumentIndexer
) -> None:
    repository = DocumentRepository(db_session)
    document = repository.create("contract.pdf", "application/pdf")

    count = indexer.index(repository, document, CONTRACT_PAGES)

    chunks = repository.list_chunks(document.id)
    assert count == 2
    assert [chunk.page_number for chunk in chunks] == [1, 2]
    assert indexer.vector_store.count() == 2

    question = indexer.embedder.embed_query("How can a party terminate the agreement?")
    best = indexer.vector_store.query(question, top_k=1)[0]
    assert best.id == chunks[1].id
    assert (best.document_id, best.page_number) == (document.id, 2)


def test_indexing_again_replaces_the_previous_index(
    db_session: Session, indexer: DocumentIndexer
) -> None:
    repository = DocumentRepository(db_session)
    document = repository.create("contract.pdf", "application/pdf")
    indexer.index(repository, document, CONTRACT_PAGES)
    old_ids = {chunk.id for chunk in repository.list_chunks(document.id)}

    indexer.index(repository, document, CONTRACT_PAGES)

    new_ids = {chunk.id for chunk in repository.list_chunks(document.id)}
    assert indexer.vector_store.count() == 2
    assert old_ids.isdisjoint(new_ids)


def test_a_document_without_text_gets_no_chunk(
    db_session: Session, indexer: DocumentIndexer
) -> None:
    repository = DocumentRepository(db_session)
    document = repository.create("blank.png", "image/png")

    assert indexer.index(repository, document, [PageText(1, "")]) == 0
    assert indexer.vector_store.count() == 0


def test_embedding_failure_keeps_the_existing_index(
    db_session: Session, make_indexer: Callable[..., DocumentIndexer]
) -> None:
    repository = DocumentRepository(db_session)
    document = repository.create("contract.pdf", "application/pdf")
    working = make_indexer(name="shared")
    working.index(repository, document, CONTRACT_PAGES)
    old_ids = [chunk.id for chunk in repository.list_chunks(document.id)]
    broken = make_indexer(unavailable=True, name="shared")

    try:
        broken.index(repository, document, CONTRACT_PAGES)
    except IndexingError as error:
        assert str(error).startswith("Indexing failed:")
        assert "connection refused" in str(error)
    else:
        raise AssertionError("IndexingError was not raised")

    assert [chunk.id for chunk in repository.list_chunks(document.id)] == old_ids
    assert working.vector_store.count() == 2


def test_remove_deletes_the_vectors_of_one_document(
    db_session: Session, indexer: DocumentIndexer
) -> None:
    repository = DocumentRepository(db_session)
    contract = repository.create("contract.pdf", "application/pdf")
    invoice = repository.create("invoice.png", "image/png")
    indexer.index(repository, contract, CONTRACT_PAGES)
    indexer.index(repository, invoice, [PageText(1, "Invoice total: 1,200 EUR")])

    indexer.remove(contract.id)

    assert indexer.vector_store.count() == 1


def test_upload_indexes_the_document(client: TestClient, indexer: DocumentIndexer) -> None:
    response = client.post(
        "/api/documents", files={"file": ("contract.docx", make_docx(), DOCX_MIME)}
    )

    assert response.status_code == 201
    assert response.json()["status"] == "ready"
    assert indexer.vector_store.count() >= 1


def test_upload_fails_cleanly_when_the_embedding_service_is_down(
    client: TestClient, make_indexer: Callable[..., DocumentIndexer]
) -> None:
    app.dependency_overrides[get_indexer] = lambda: make_indexer(unavailable=True, name="down")

    response = client.post(
        "/api/documents", files={"file": ("contract.docx", make_docx(), DOCX_MIME)}
    )

    body = response.json()
    assert response.status_code == 201
    assert body["status"] == "failed"
    assert body["error"].startswith("Indexing failed:")
    pages = client.get(f"/api/documents/{body['id']}/pages")
    assert pages.json() == []