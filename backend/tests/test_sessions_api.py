"""Tests for the /api/sessions routes (fake embeddings, fake LLM, real Chroma)."""

import io
from collections.abc import Callable, Iterator

import docx
import pytest
from fastapi.testclient import TestClient

from app.api.dependencies import get_indexer
from app.core.config import get_settings
from app.db.repositories import DocumentRepository
from app.main import app
from app.rag.answering import REFUSAL_MESSAGE
from app.retrieval.indexing import DocumentIndexer

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
QUESTION = "How can a party terminate the agreement?"
REPLY = "Either party may terminate with 30 days notice [S1]."


def make_docx() -> bytes:
    """Return a Word document with two heading-based sections."""
    document = docx.Document()
    document.add_heading("Start date", level=1)
    document.add_paragraph("This service agreement starts on 1 January 2026 in Paris.")
    document.add_heading("Termination", level=1)
    document.add_paragraph("Either party may terminate the agreement with 30 days notice.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(client: TestClient, filename: str = "contract.docx") -> str:
    """Upload a ready document and return its id."""
    response = client.post("/api/documents", files={"file": (filename, make_docx(), DOCX_MIME)})
    assert response.json()["status"] == "ready"
    return str(response.json()["id"])


def new_session(client: TestClient) -> str:
    response = client.post("/api/sessions")
    assert response.status_code == 201
    return str(response.json()["id"])


@pytest.fixture
def set_min_score(client: TestClient) -> Iterator[Callable[[float], None]]:
    """Change the relevance threshold for one test (see test_ask_api.py)."""

    def apply(value: float) -> None:
        settings = get_settings().model_copy(update={"retrieval_min_score": value})
        app.dependency_overrides[get_settings] = lambda: settings

    yield apply
    app.dependency_overrides.pop(get_settings, None)


def test_create_returns_an_empty_session(client: TestClient) -> None:
    response = client.post("/api/sessions")

    assert response.status_code == 201
    body = response.json()
    assert len(body["id"]) == 36
    assert body["title"] is None
    assert body["message_count"] == 0
    assert body["created_at"].endswith("Z") or "+00:00" in body["created_at"]


def test_list_is_empty_on_a_new_database(client: TestClient) -> None:
    assert client.get("/api/sessions").json() == []


def test_read_returns_an_empty_new_session(client: TestClient) -> None:
    session_id = new_session(client)

    body = client.get(f"/api/sessions/{session_id}").json()

    assert body["id"] == session_id
    assert body["messages"] == []
    assert body["document_ids"] == []


def test_unknown_session_gives_404(client: TestClient, fake_llm) -> None:
    assert client.get("/api/sessions/does-not-exist").status_code == 404

    response = client.post("/api/sessions/does-not-exist/ask", json={"question": QUESTION})

    assert response.status_code == 404
    assert response.json()["detail"] == "Session not found."
    assert fake_llm.calls == []


def test_ask_in_session_saves_the_question_and_the_cited_answer(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.0)
    document_id = upload(client)
    session_id = new_session(client)
    fake_llm.replies.append(REPLY)

    response = client.post(
        f"/api/sessions/{session_id}/ask",
        json={"question": QUESTION, "top_k": 1, "document_ids": [document_id]},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["session"]["title"] == QUESTION
    assert body["session"]["message_count"] == 2
    assert body["question"]["role"] == "user"
    assert body["question"]["content"] == QUESTION
    assert body["question"]["status"] is None
    answer = body["answer"]
    assert answer["role"] == "assistant"
    assert answer["mode"] == "ask"
    assert answer["status"] == "answered"
    assert answer["refusal_reason"] is None
    citation = answer["citations"][0]
    assert citation["filename"] == "contract.docx"
    expected = f"Either party may terminate with 30 days notice {citation['display']}."
    assert answer["content"] == expected
    assert answer["agent_trace"] == []


def test_a_saved_session_reads_back_exactly_like_the_live_answer(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    """Closing the tab and coming back gives the same messages and citations."""
    set_min_score(0.0)
    document_id = upload(client)
    session_id = new_session(client)
    fake_llm.replies.append(REPLY)
    live = client.post(
        f"/api/sessions/{session_id}/ask",
        json={"question": QUESTION, "top_k": 1, "document_ids": [document_id]},
    ).json()

    restored = client.get(f"/api/sessions/{session_id}").json()

    assert restored["title"] == QUESTION
    assert restored["messages"] == [live["question"], live["answer"]]
    assert restored["document_ids"] == [document_id]


def test_ask_in_session_gives_the_same_answer_as_api_ask(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.0)
    upload(client)
    session_id = new_session(client)
    fake_llm.replies.extend([REPLY, REPLY])

    direct = client.post("/api/ask", json={"question": QUESTION, "top_k": 1}).json()
    saved = client.post(
        f"/api/sessions/{session_id}/ask", json={"question": QUESTION, "top_k": 1}
    ).json()["answer"]

    assert saved["content"] == direct["answer"]
    assert saved["citations"] == direct["citations"]
    assert saved["best_score"] == direct["best_score"]


def test_a_refusal_is_saved_with_its_reason(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.999)
    upload(client)
    session_id = new_session(client)

    client.post(f"/api/sessions/{session_id}/ask", json={"question": "What is the CEO's salary?"})

    answer = client.get(f"/api/sessions/{session_id}").json()["messages"][1]
    assert answer["status"] == "refused"
    assert answer["refusal_reason"] == "no_relevant_passages"
    assert answer["content"] == REFUSAL_MESSAGE
    assert answer["citations"] == []
    assert fake_llm.calls == []


def test_nothing_is_saved_when_the_model_is_down(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.0)
    upload(client)
    session_id = new_session(client)
    fake_llm.unavailable = True

    response = client.post(f"/api/sessions/{session_id}/ask", json={"question": QUESTION})

    assert response.status_code == 503
    assert response.json()["detail"].startswith("Answer generation failed")
    body = client.get(f"/api/sessions/{session_id}").json()
    assert body["messages"] == []
    assert body["title"] is None


def test_nothing_is_saved_when_the_embedding_service_is_down(
    client: TestClient, make_indexer: Callable[..., DocumentIndexer]
) -> None:
    session_id = new_session(client)
    app.dependency_overrides[get_indexer] = lambda: make_indexer(unavailable=True, name="down")

    response = client.post(f"/api/sessions/{session_id}/ask", json={"question": QUESTION})

    assert response.status_code == 503
    assert "Search failed" in response.json()["detail"]
    assert client.get(f"/api/sessions/{session_id}").json()["messages"] == []


def test_the_documents_of_the_latest_question_are_kept(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.999)
    contract_id = upload(client, "contract.docx")
    other_id = upload(client, "other.docx")
    session_id = new_session(client)
    url = f"/api/sessions/{session_id}/ask"

    client.post(url, json={"question": "Q1?", "document_ids": [contract_id, other_id]})
    client.post(url, json={"question": "Q2?", "document_ids": [other_id, "unknown-id"]})

    assert client.get(f"/api/sessions/{session_id}").json()["document_ids"] == [other_id]


def test_no_document_ids_records_every_ready_document(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.999)
    first = upload(client, "contract.docx")
    second = upload(client, "other.docx")
    session_id = new_session(client)

    client.post(f"/api/sessions/{session_id}/ask", json={"question": "Q?"})

    document_ids = client.get(f"/api/sessions/{session_id}").json()["document_ids"]
    assert sorted(document_ids) == sorted([first, second])


def test_citations_stay_readable_after_the_document_is_reindexed(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    """Re-indexing replaces the chunks: the saved copy of the source is still shown."""
    set_min_score(0.0)
    document_id = upload(client)
    session_id = new_session(client)
    fake_llm.replies.append(REPLY)
    live = client.post(
        f"/api/sessions/{session_id}/ask", json={"question": QUESTION, "top_k": 1}
    ).json()
    with app.state.session_factory() as db:
        repository = DocumentRepository(db)
        document = repository.get(document_id)
        assert document is not None
        repository.replace_chunks(document, [])

    restored = client.get(f"/api/sessions/{session_id}").json()["messages"][1]

    assert restored["citations"] == live["answer"]["citations"]
    assert "terminate" in restored["citations"][0]["text"]


def test_the_session_list_shows_titles_and_counts_most_recent_first(
    client: TestClient, fake_llm, set_min_score: Callable[[float], None]
) -> None:
    set_min_score(0.999)
    upload(client)
    older = new_session(client)
    newer = new_session(client)
    empty = new_session(client)
    client.post(f"/api/sessions/{older}/ask", json={"question": "Older question?"})
    client.post(f"/api/sessions/{newer}/ask", json={"question": "Newer question?"})
    client.post(f"/api/sessions/{newer}/ask", json={"question": "Follow-up?"})

    summaries = client.get("/api/sessions").json()

    assert [(s["id"], s["title"], s["message_count"]) for s in summaries] == [
        (newer, "Newer question?", 4),
        (older, "Older question?", 2),
        (empty, None, 0),
    ]


@pytest.mark.parametrize(
    "payload",
    [
        {},
        {"question": "   "},
        {"question": "x" * 1001},
        {"question": "ok", "top_k": 0},
        {"question": "ok", "top_k": 21},
    ],
)
def test_ask_in_session_refuses_invalid_requests(client: TestClient, payload: dict) -> None:
    session_id = new_session(client)

    assert client.post(f"/api/sessions/{session_id}/ask", json=payload).status_code == 422