"""Create the demo quarterly business review: a fictional PPTX presentation.

The review covers Q2 2026 (April to June) of the demo contract
(service_agreement.pdf, Agreement No. HDS-2026-014) and repeats the demo
invoice (HDS-INV-2026-0042, USD 15,950.00). Its figures follow the contract:
availability target of 99.5% and service credits (Article 2.6), fees and
rates (Article 3.1), price review after twelve months (Article 3.6).

ATLAS reads one slide as one page, so each slide holds one topic and every
figure has a single page to cite. Two slides hold a table, which the PPTX
parser reads row by row.

Usage, from the backend folder:
    python -m scripts.make_demo_review

The file is written to demo_documents/ at the root of the repository, then
read back with the ATLAS parser to check that the key values are found on
the right slides. Running the script again gives exactly the same file.
"""

import io
import math
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation as PresentationFile
from pptx.slide import Slide
from pptx.util import Emu, Inches, Pt

from app.ingestion.models import ParseOptions
from app.ingestion.pptx_parser import parse_pptx

OUTPUT = (
    Path(__file__).resolve().parents[2] / "demo_documents" / "quarterly_review_HDS_Q2_2026.pptx"
)

# Fixed date written in the file (properties and zip entries), so that
# running the script again gives a byte-for-byte identical file.
FIXED_DATE = datetime(2026, 7, 8, 9, 0, 0)

# --- Figures of the review ----------------------------------------------------

AVAILABILITY_TARGET = Decimal("99.5")  # Article 2.5
MONTHLY_FEE = Decimal(12500)  # Article 3.1
LICENCE_PRICE = Decimal(40)  # per additional user per month, Article 3.1
CONSULTING_DAY_RATE = Decimal(950)  # Article 3.1
CREDIT_PER_STEP = Decimal("0.05")  # 5% of the Monthly Fee per 0.5% below target
CREDIT_STEP = Decimal("0.5")
CREDIT_CAP = Decimal("0.25")  # at most 25% of the Monthly Fee


@dataclass(frozen=True)
class MonthOfService:
    month: str
    availability: Decimal  # percentage of the month
    critical_incidents: int
    max_response: str  # slowest response to a critical incident, "–" if none
    hours_in_month: int


@dataclass(frozen=True)
class Invoice:
    number: str
    date: str
    licences: int  # additional user licences of the previous month (in arrears)
    consulting_days: int  # extra consulting days of the previous month (in arrears)

    @property
    def extras(self) -> Decimal:
        return self.licences * LICENCE_PRICE + self.consulting_days * CONSULTING_DAY_RATE

    @property
    def total(self) -> Decimal:
        return MONTHLY_FEE + self.extras


MONTHS = [
    MonthOfService("April 2026", Decimal("99.91"), 0, "–", 30 * 24),
    MonthOfService("May 2026", Decimal("99.38"), 2, "1 h 25 min", 31 * 24),
    MonthOfService("June 2026", Decimal("99.94"), 0, "–", 30 * 24),
]

# The April invoice is the demo invoice: 15 licences and 3 consulting days of March.
INVOICES = [
    Invoice("HDS-INV-2026-0042", "1 April 2026", licences=15, consulting_days=3),
    Invoice("HDS-INV-2026-0057", "1 May 2026", licences=27, consulting_days=2),
    Invoice("HDS-INV-2026-0071", "1 June 2026", licences=36, consulting_days=0),
]


def money(amount: Decimal) -> str:
    """12500 -> '12,500.00', as on the demo invoice."""
    return f"{amount:,.2f}"


def average_availability() -> Decimal:
    """Average of the three months, rounded to two decimals."""
    total = sum((month.availability for month in MONTHS), Decimal(0))
    return (total / len(MONTHS)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def downtime(month: MonthOfService) -> str:
    """Time the platform was unavailable, e.g. '4 h 37 min'."""
    minutes = month.hours_in_month * 60 * (100 - month.availability) / 100
    rounded = int(minutes.quantize(Decimal(1), rounding=ROUND_HALF_UP))
    return f"{rounded // 60} h {rounded % 60} min"


def service_credit(month: MonthOfService) -> Decimal:
    """Article 2.6: 5% of the Monthly Fee for each 0.5% (or part of it) below target."""
    shortfall = AVAILABILITY_TARGET - month.availability
    if shortfall <= 0:
        return Decimal(0)
    steps = math.ceil(shortfall / CREDIT_STEP)
    rate = min(steps * CREDIT_PER_STEP, CREDIT_CAP)
    return MONTHLY_FEE * rate


def quarter_total() -> Decimal:
    return sum((invoice.total for invoice in INVOICES), Decimal(0))


MAY = MONTHS[1]

# Values read back from each slide after saving (slide number -> texts).
EXPECTED_ON_SLIDE = {
    1: ["Quarterly Business Review", "Q2 2026", "HDS-2026-014"],
    2: ["24 April 2026", "99.74%", "USD 45,370.00"],
    3: ["May 2026 | 99.38% | 99.5% | 2 | 1 h 25 min"],
    4: ["205,120"],
    5: ["HDS-INV-2026-0042 | 1 April 2026 | 12,500.00 | 3,450.00 | 15,950.00"],
    6: ["4 h 37 min", "USD 625.00"],
}

# --- Look of the slides --------------------------------------------------------

TEAL = RGBColor(0x0C, 0x7C, 0x7A)
CORAL = RGBColor(0xFF, 0x7F, 0x66)
INK = RGBColor(0x16, 0x32, 0x3A)
MUTED = RGBColor(0x5B, 0x72, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_TINT = RGBColor(0xE8, 0xF4, 0xF3)
FONT = "Arial"  # available on every Mac and Windows machine

SLIDE_WIDTH = Inches(13.333)  # 16:9
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_WIDTH = Emu(SLIDE_WIDTH - 2 * MARGIN)
BLANK_LAYOUT = 6  # the layout without placeholders in the default template

FOOTER = "Quarterly Business Review Q2 2026 - Orion Freight Co. - Helix Data Services Ltd."
FICTION_NOTE = "Fictional document created for demonstration purposes only."


def add_text(
    slide: Slide,
    top: float,
    height: float,
    lines: list[str],
    size: int,
    color: RGBColor = INK,
    bold: bool = False,
    space_after: int = 0,
) -> None:
    """A text box across the slide, one paragraph per line (top and height in inches)."""
    box = slide.shapes.add_textbox(MARGIN, Inches(top), CONTENT_WIDTH, Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_bullets(slide: Slide, top: float, height: float, items: list[str]) -> None:
    add_text(slide, top, height, [f"•  {item}" for item in items], size=20, space_after=14)


def add_table(slide: Slide, top: float, rows: list[list[str]]) -> None:
    """A table with a teal header row and tinted body rows (top in inches)."""
    row_height = Inches(0.55)
    shape = slide.shapes.add_table(
        len(rows), len(rows[0]), MARGIN, Inches(top), CONTENT_WIDTH, Emu(row_height * len(rows))
    )
    table = shape.table
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = TEAL
            else:
                cell.fill.fore_color.rgb = ROW_TINT if row_index % 2 else WHITE
            paragraph = cell.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(16)
            run.font.bold = row_index == 0
            run.font.color.rgb = WHITE if row_index == 0 else INK


def new_slide(presentation: PresentationFile, title: str, number: int) -> Slide:
    """A blank slide with its title, a coral accent line and the footer."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])
    add_text(slide, 0.5, 0.9, [title], size=32, color=TEAL, bold=True)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.4), Inches(1.2), Inches(0.06)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = CORAL
    accent.line.fill.background()
    add_text(slide, 6.9, 0.4, [f"{FOOTER} - {number} / 6"], size=11, color=MUTED)
    return slide


def build_presentation() -> PresentationFile:
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    # 1. Title
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])
    add_text(slide, 2.0, 1.0, ["Quarterly Business Review - Q2 2026"], 40, TEAL, bold=True)
    add_text(
        slide,
        3.1,
        1.4,
        [
            "Orion Freight Co. account - April to June 2026",
            "Helix Data Services Ltd. - Agreement No. HDS-2026-014",
        ],
        size=22,
        space_after=6,
    )
    add_text(
        slide,
        4.8,
        0.9,
        ["Prepared on 8 July 2026 by Priya Nair, Account Director, Helix Data Services Ltd."],
        size=16,
        color=MUTED,
    )
    add_text(slide, 6.9, 0.4, [FICTION_NOTE], size=11, color=MUTED)

    # 2. Executive summary
    slide = new_slide(presentation, "Executive summary", 2)
    add_bullets(
        slide,
        1.8,
        4.8,
        [
            ("Onboarding completed on 24 April 2026, within the eight-week period of Article 2.2."),
            (
                f"Average platform availability over the quarter: {average_availability()}%. "
                f"May 2026 was below the {AVAILABILITY_TARGET}% target."
            ),
            "Two critical incidents, both answered within two hours.",
            f"Total invoiced in Q2 2026: USD {money(quarter_total())}.",
            "Next quarterly review: October 2026.",
        ],
    )

    # 3. Service levels
    slide = new_slide(presentation, "Service levels", 3)
    rows = [["Month", "Availability", "Target", "Critical incidents", "Max response time"]]
    for month in MONTHS:
        rows.append(
            [
                month.month,
                f"{month.availability}%",
                f"{AVAILABILITY_TARGET}%",
                str(month.critical_incidents),
                month.max_response,
            ]
        )
    add_table(slide, 1.9, rows)
    add_text(
        slide,
        4.4,
        1.2,
        [
            (
                "Availability excludes planned maintenance announced at least five Business "
                "Days in advance (Article 2.5)."
            )
        ],
        size=16,
        color=MUTED,
    )

    # 4. Usage and operations
    slide = new_slide(presentation, "Usage and operations", 4)
    add_bullets(
        slide,
        1.8,
        4.8,
        [
            "Active platform users: 42 in April, 51 in May and 58 in June 2026.",
            "Shipments processed: 184,300 in April, 197,650 in May and 205,120 in June 2026.",
            "Monthly operations reports delivered on time: 3 of 3.",
            "Support tickets opened in the quarter: 37, of which 35 are resolved.",
        ],
    )

    # 5. Fees invoiced
    slide = new_slide(presentation, "Fees invoiced in Q2 2026 (USD)", 5)
    rows = [["Invoice", "Date", "Monthly Fee", "Extras", "Total"]]
    for invoice in INVOICES:
        rows.append(
            [
                invoice.number,
                invoice.date,
                money(MONTHLY_FEE),
                money(invoice.extras),
                money(invoice.total),
            ]
        )
    add_table(slide, 1.9, rows)
    add_text(
        slide,
        4.4,
        2.2,
        [
            (
                f"Total invoiced in the quarter: USD {money(quarter_total())}. "
                "All invoices were paid within 30 days."
            ),
            (
                "Extras are additional user licences (USD 40 per user per month) and extra "
                "consulting days (USD 950 per day) of the previous month, invoiced in arrears."
            ),
        ],
        size=16,
        space_after=10,
    )

    # 6. Risks and next steps
    slide = new_slide(presentation, "Risks and next steps", 6)
    add_bullets(
        slide,
        1.8,
        4.8,
        [
            (
                f"On 14 May 2026, a storage failure in the primary region (Frankfurt) made "
                f"the platform unavailable for {downtime(MAY)}."
            ),
            (
                f"May availability ({MAY.availability}%) gives a service credit of "
                f"USD {money(service_credit(MAY))} (5% of the Monthly Fee, Article 2.6), "
                "deducted from the July 2026 invoice."
            ),
            "A price review is not possible before March 2027 (Article 3.6).",
            "Next quarterly review: October 2026.",
        ],
    )
    return presentation


def set_fixed_properties(presentation: PresentationFile) -> None:
    properties = presentation.core_properties
    properties.title = "Quarterly Business Review - Q2 2026"
    properties.author = "Helix Data Services Ltd. (fictional)"
    properties.last_modified_by = "ATLAS demo script"
    properties.created = FIXED_DATE
    properties.modified = FIXED_DATE
    properties.last_printed = FIXED_DATE
    properties.revision = 1


def to_bytes() -> bytes:
    """The PPTX file, with fixed dates inside the zip so the bytes never change."""
    presentation = build_presentation()
    set_fixed_properties(presentation)
    saved = io.BytesIO()
    presentation.save(saved)

    # python-pptx stamps every zip entry with the current time: rewrite them.
    fixed = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(saved.getvalue())) as source,
        zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for entry in source.infolist():
            stamped = zipfile.ZipInfo(entry.filename, date_time=FIXED_DATE.timetuple()[:6])
            stamped.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(stamped, source.read(entry.filename))
    return fixed.getvalue()


def missing_values(path: Path) -> list[str]:
    """Read the file with the ATLAS parser; list the expected texts not found."""
    pages = {page.page_number: page.text for page in parse_pptx(path, ParseOptions())}
    missing = []
    for number, values in EXPECTED_ON_SLIDE.items():
        for value in values:
            if value not in pages.get(number, ""):
                missing.append(f"slide {number}: {value}")
    if len(pages) != len(EXPECTED_ON_SLIDE):
        missing.append(f"{len(EXPECTED_ON_SLIDE)} slides expected, {len(pages)} found")
    return missing


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_bytes(to_bytes())
    print(f"Created {OUTPUT}")

    missing = missing_values(OUTPUT)
    if missing:
        raise SystemExit("Parser check failed, not found:\n" + "\n".join(missing))
    print(f"Parser check passed: {len(EXPECTED_ON_SLIDE)} slides, key values on the right slide")


if __name__ == "__main__":
    main()