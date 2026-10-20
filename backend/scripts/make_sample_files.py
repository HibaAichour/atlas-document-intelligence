"""Create small fictional DOCX and PPTX files to test document ingestion.

Usage (from the repository root, with the backend virtual environment active):
    python backend/scripts/make_sample_files.py

The files are written to your home folder, outside the repository.
"""

from pathlib import Path

from docx import Document
from pptx import Presentation

OUTPUT_DIR = Path.home()


def make_docx(path: Path) -> None:
    """Three sections, each starting with a heading."""
    document = Document()
    document.add_heading("Service Agreement Summary", level=1)
    document.add_paragraph(
        "This fictional agreement is between Northwind Consulting Ltd. "
        "and Contoso Analytics Inc."
    )
    document.add_heading("Termination", level=1)
    document.add_paragraph(
        "Either party may terminate this agreement with 60 days written notice."
    )
    document.add_heading("Fees", level=1)
    document.add_paragraph(
        "The monthly fee is 4,500.00 EUR, payable within 30 days of invoice."
    )
    document.save(path)


def make_pptx(path: Path) -> None:
    """Two slides with a title and a bulleted body."""
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # "Title and Content"

    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Q3 Business Review"
    slide.placeholders[1].text = "Revenue: 1.2M EUR (+8% vs Q2)\nNew clients: 14"

    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Q4 Priorities"
    slide.placeholders[1].text = "Launch analytics dashboard\nHire two data engineers"

    presentation.save(path)


if __name__ == "__main__":
    docx_path = OUTPUT_DIR / "test_document.docx"
    pptx_path = OUTPUT_DIR / "test_presentation.pptx"
    make_docx(docx_path)
    make_pptx(pptx_path)
    print(f"Created {docx_path}")
    print(f"Created {pptx_path}")