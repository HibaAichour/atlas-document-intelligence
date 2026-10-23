"""Lay out an ExportContent as a PowerPoint presentation (.pptx), built in memory.

Slides: a title slide, one slide with the answer (or the table of fields)
and the Critic's verdict, then the sources, two per slide. A slide has
little room, so long passages are shortened; the Word export keeps them whole.
"""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation as PresentationFile
from pptx.slide import Slide
from pptx.table import _Cell
from pptx.util import Emu, Inches, Pt

from app.export.content import ExportContent, ExportReview, ExportSource

TEAL = RGBColor(0x0C, 0x7C, 0x7A)
CORAL = RGBColor(0xFF, 0x7F, 0x66)
CORAL_TEXT = RGBColor(0xD9, 0x5A, 0x43)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
INK = RGBColor(0x16, 0x32, 0x3A)
MUTED = RGBColor(0x5B, 0x72, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_TINT = RGBColor(0xE8, 0xF4, 0xF3)
PASSAGE_FILL = RGBColor(0xF2, 0xF8, 0xF7)
FONT = "Arial"

SLIDE_WIDTH = Inches(13.333)  # 16:9
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_WIDTH = Emu(SLIDE_WIDTH - 2 * MARGIN)
BLANK_LAYOUT = 6

SOURCES_PER_SLIDE = 2
MAX_PASSAGE_CHARS = 700  # beyond this, a passage is cut at a word and ends with "…"
MAX_ISSUES_ON_SLIDE = 4


def render_pptx(content: ExportContent) -> bytes:
    """The .pptx file of the content, as bytes."""
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    properties = presentation.core_properties
    properties.title = content.title
    properties.author = content.texts.app_name
    properties.created = content.generated_at
    properties.modified = content.generated_at
    properties.subject = content.session_title or ""

    title_slide(presentation, content)
    if content.kind == "answer":
        answer_slide(presentation, content)
    else:
        fields_slide(presentation, content)

    groups = [
        content.sources[start : start + SOURCES_PER_SLIDE]
        for start in range(0, len(content.sources), SOURCES_PER_SLIDE)
    ]
    for index, group in enumerate(groups, start=1):
        sources_slide(presentation, content, group, index, len(groups))

    saved = io.BytesIO()
    presentation.save(saved)
    return saved.getvalue()


def shorten(text: str, limit: int = MAX_PASSAGE_CHARS) -> str:
    """Cut a long text at the last space before the limit and add "…"."""
    if len(text) <= limit:
        return text
    cut = text[:limit].rsplit(" ", 1)[0].rstrip(" ,;:.")
    return cut + " …"


def add_text(
    slide: Slide,
    left: int,
    top: int,
    width: int,
    height: int,
    paragraphs: list[tuple[str, int, RGBColor, bool]],
    space_after: int = 6,
) -> None:
    """A text box; each paragraph is (text, size, colour, bold)."""
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, (text, size, color, bold) in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def new_slide(presentation: PresentationFile, title: str, content: ExportContent) -> Slide:
    """A blank slide with its title, a coral accent line and the footer."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])
    add_text(slide, MARGIN, Inches(0.45), CONTENT_WIDTH, Inches(0.9), [(title, 30, TEAL, True)])
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.3), Inches(1.2), Inches(0.06)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = CORAL
    accent.line.fill.background()
    footer(slide, content)
    return slide


def footer(slide: Slide, content: ExportContent) -> None:
    texts = content.texts
    text = texts.footer.format(date=texts.format_date(content.generated_at))
    add_text(slide, MARGIN, Inches(6.95), CONTENT_WIDTH, Inches(0.4), [(text, 10, MUTED, False)])


def title_slide(presentation: PresentationFile, content: ExportContent) -> None:
    texts = content.texts
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])
    add_text(
        slide,
        MARGIN,
        Inches(1.6),
        CONTENT_WIDTH,
        Inches(2.2),
        [(texts.app_name, 18, MUTED, True), (content.title, 36, TEAL, True)],
    )
    if content.request:
        add_text(
            slide,
            MARGIN,
            Inches(4.2),
            CONTENT_WIDTH,
            Inches(1.6),
            [(content.request, 22, INK, False)],
        )
    footer(slide, content)


def answer_slide(presentation: PresentationFile, content: ExportContent) -> None:
    texts = content.texts
    slide = new_slide(presentation, texts.answer_heading, content)
    answer = content.answer or ""
    size = 22 if len(answer) <= 300 else 18 if len(answer) <= 600 else 15
    paragraphs: list[tuple[str, int, RGBColor, bool]] = []
    if content.request:
        paragraphs.append((content.request, 16, MUTED, False))
    paragraphs.append((answer, size, INK, False))
    add_text(slide, MARGIN, Inches(1.6), CONTENT_WIDTH, Inches(3.0), paragraphs, space_after=12)
    if content.review is not None:
        review_box(slide, content.review, content, top=Inches(4.75))


def fields_slide(presentation: PresentationFile, content: ExportContent) -> None:
    texts = content.texts
    slide = new_slide(presentation, content.title, content)
    rows = content.rows
    row_height = Inches(0.42)
    shape = slide.shapes.add_table(
        len(rows) + 1, 3, MARGIN, Inches(1.6), CONTENT_WIDTH, Emu(row_height * (len(rows) + 1))
    )
    table = shape.table
    widths = (Inches(3.6), Inches(6.6), Inches(1.733))
    for index, width in enumerate(widths):
        table.columns[index].width = Emu(width)

    header = (texts.field_column, texts.value_column, texts.page_column)
    for column, title in enumerate(header):
        fill_cell(table.cell(0, column), [(title, 14, WHITE, True)], TEAL)
    for index, row in enumerate(rows, start=1):
        tint = ROW_TINT if index % 2 else WHITE
        value_color = MUTED if row.value == texts.not_found else INK
        value_lines = [(row.value, 14, value_color, False)]
        value_lines += [(problem, 11, CORAL_TEXT, False) for problem in row.issues]
        fill_cell(table.cell(index, 0), [(row.label, 14, INK, False)], tint)
        fill_cell(table.cell(index, 1), value_lines, tint)
        fill_cell(table.cell(index, 2), [(row.page, 14, INK, False)], tint)

    if content.review is not None:
        top = Inches(1.6) + row_height * (len(rows) + 1) + Inches(0.3)
        review_box(slide, content.review, content, top=Emu(min(top, Inches(6.0))))


def fill_cell(cell: _Cell, lines: list[tuple[str, int, RGBColor, bool]], fill: RGBColor) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    frame = cell.text_frame
    for index, (text, size, color, bold) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def review_box(slide: Slide, review: ExportReview, content: ExportContent, top: int) -> None:
    """The verdict: one line, then the problems (the first few on a slide)."""
    texts = content.texts
    ok = not review.certain and not review.to_check and review.note is None
    lines: list[tuple[str, int, RGBColor, bool]] = [
        (review.summary, 16, GREEN if ok else CORAL_TEXT, True)
    ]
    problems = [(texts.review_certain, problem) for problem in review.certain]
    problems += [(texts.review_to_check, problem) for problem in review.to_check]
    for group, problem in problems[:MAX_ISSUES_ON_SLIDE]:
        lines.append((f"{group}: {problem}", 13, INK, False))
    hidden = len(problems) - MAX_ISSUES_ON_SLIDE
    if hidden > 0:
        lines.append((f"+ {hidden}", 13, MUTED, False))
    if review.note:
        lines.append((review.note, 12, MUTED, False))
    height = Emu(max(Inches(6.85) - top, Inches(0.6)))
    add_text(slide, MARGIN, top, CONTENT_WIDTH, height, lines, space_after=4)


def sources_slide(
    presentation: PresentationFile,
    content: ExportContent,
    sources: list[ExportSource],
    number: int,
    total: int,
) -> None:
    texts = content.texts
    title = texts.sources_heading if total == 1 else f"{texts.sources_heading} ({number}/{total})"
    slide = new_slide(presentation, title, content)
    gap = Inches(0.25)
    height = Emu((Inches(5.2) - gap * (SOURCES_PER_SLIDE - 1)) // SOURCES_PER_SLIDE)
    for index, source in enumerate(sources):
        top = Emu(Inches(1.6) + index * (height + gap))
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, top, CONTENT_WIDTH, height)
        panel.fill.solid()
        panel.fill.fore_color.rgb = PASSAGE_FILL
        panel.line.fill.background()
        panel.shadow.inherit = False
        lines: list[tuple[str, int, RGBColor, bool]] = [(source.display, 15, TEAL, True)]
        if source.used_for:
            labels = ", ".join(source.used_for)
            lines.append((texts.source_used_for.format(labels=labels), 12, MUTED, False))
        # The passage keeps its line breaks (one paragraph per line of the page).
        for passage_line in shorten(source.text).split("\n"):
            if passage_line.strip():
                lines.append((passage_line, 13, INK, False))
        inset = Inches(0.2)
        add_text(
            slide,
            MARGIN + inset,
            top + Inches(0.1),
            CONTENT_WIDTH - 2 * inset,
            height - Inches(0.2),
            lines,
            space_after=4,
        )