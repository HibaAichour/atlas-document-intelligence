"""Lay out an ExportContent as a Word document (.docx), built in memory.

The Word file is the complete version: every source passage is printed in
full. Colours follow the interface: teal headings, coral for problems.
"""

import io
from typing import cast

from docx import Document
from docx.document import Document as DocumentFile
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from app.export.content import ExportContent, ExportReview

TEAL = RGBColor(0x0C, 0x7C, 0x7A)
CORAL = RGBColor(0xD9, 0x5A, 0x43)  # darker than the interface coral, readable on white
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
INK = RGBColor(0x16, 0x32, 0x3A)
MUTED = RGBColor(0x5B, 0x72, 0x78)
PASSAGE_SHADE = "EEF6F5"  # light teal behind the source passages
HEADER_SHADE = "0C7C7A"
FONT = "Arial"


def render_docx(content: ExportContent) -> bytes:
    """The .docx file of the content, as bytes."""
    texts = content.texts
    document = Document()
    set_base_style(document)
    properties = document.core_properties
    properties.title = content.title
    properties.author = texts.app_name
    properties.created = content.generated_at
    properties.modified = content.generated_at
    properties.subject = content.session_title or ""

    heading(document, content.title, size=22)

    if content.request:
        heading(
            document,
            texts.question_heading if content.kind == "answer" else texts.request_heading,
        )
        line(document, content.request)

    if content.answer is not None:
        heading(document, texts.answer_heading)
        line(document, content.answer, size=12)
    if content.rows:
        heading(document, texts.fields_heading)
        fields_table(document, content)

    if content.review is not None:
        heading(document, texts.review_heading)
        review_block(document, content.review, content)

    if content.sources:
        heading(document, texts.sources_heading)
        for source in content.sources:
            line(document, source.display, color=TEAL, bold=True, space_after=2)
            if source.used_for:
                labels = ", ".join(source.used_for)
                line(
                    document,
                    texts.source_used_for.format(labels=labels),
                    color=MUTED,
                    size=9,
                    italic=True,
                    space_after=2,
                )
            for paragraph_text in source.text.split("\n"):
                if paragraph_text.strip():
                    passage = line(document, paragraph_text, size=10, space_after=0)
                    shade(passage, PASSAGE_SHADE)
                    passage.paragraph_format.left_indent = Cm(0.3)
            document.add_paragraph().paragraph_format.space_after = Pt(4)

    footer = document.sections[0].footer.paragraphs[0]
    run = footer.add_run(texts.footer.format(date=texts.format_date(content.generated_at)))
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    saved = io.BytesIO()
    document.save(saved)
    return saved.getvalue()


def set_base_style(document: DocumentFile) -> None:
    style = cast(ParagraphStyle, document.styles["Normal"])
    style.font.name = FONT
    style.font.size = Pt(11)
    style.font.color.rgb = INK
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)


def heading(document: DocumentFile, text: str, size: int = 14) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14 if size < 20 else 0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = TEAL
    return paragraph


def line(
    document: DocumentFile,
    text: str,
    *,
    color: RGBColor = INK,
    size: int = 11,
    bold: bool = False,
    italic: bool = False,
    space_after: int = 6,
) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.font.color.rgb = color
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return paragraph


def bullet(document: DocumentFile, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.size = Pt(10)


def shade(paragraph: Paragraph, color: str) -> None:
    """A background colour behind a paragraph (no python-docx API for it)."""
    properties = paragraph._p.get_or_add_pPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def shade_cell(cell: _Cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def fields_table(document: DocumentFile, content: ExportContent) -> None:
    texts = content.texts
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = (Cm(5.0), Cm(9.0), Cm(2.5))
    header = table.rows[0].cells
    for cell, title, width in zip(
        header, (texts.field_column, texts.value_column, texts.page_column), widths, strict=True
    ):
        cell.width = width
        shade_cell(cell, HEADER_SHADE)
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in content.rows:
        cells = table.add_row().cells
        for cell, width in zip(cells, widths, strict=True):
            cell.width = width
        cells[0].paragraphs[0].add_run(row.label)
        value = cells[1].paragraphs[0].add_run(row.value)
        if row.value == texts.not_found:
            value.italic = True
            value.font.color.rgb = MUTED
        for problem in row.issues:
            run = cells[1].add_paragraph().add_run(problem)
            run.font.size = Pt(9)
            run.font.color.rgb = CORAL
        cells[2].paragraphs[0].add_run(row.page)
    for table_row in table.rows:
        for cell in table_row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)


def review_block(document: DocumentFile, review: ExportReview, content: ExportContent) -> None:
    texts = content.texts
    ok = not review.certain and not review.to_check and review.note is None
    line(document, review.summary, color=GREEN if ok else CORAL, bold=True)
    if review.certain:
        line(document, texts.review_certain, bold=True, size=10, space_after=2)
        for problem in review.certain:
            bullet(document, problem)
    if review.to_check:
        line(document, texts.review_to_check, bold=True, size=10, space_after=2)
        line(document, texts.review_to_check_note, color=MUTED, size=9, italic=True, space_after=2)
        for problem in review.to_check:
            bullet(document, problem)
    if review.note:
        line(document, review.note, color=MUTED, size=9, italic=True)