"""Read PPTX files: one slide is one page."""

from collections.abc import Iterable, Iterator
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape

from app.ingestion.models import PageText, ParseError, ParseOptions
from app.ingestion.text import clean_text


def parse_pptx(path: Path, options: ParseOptions) -> list[PageText]:
    """Return the text of every slide, in order. Empty slides keep their number."""
    try:
        presentation = Presentation(str(path))
    except Exception as error:  # A damaged file can fail in many different ways.
        raise ParseError(f"Cannot open the PPTX file ({error}).") from error

    slides = list(presentation.slides)
    if len(slides) > options.max_pages:
        raise ParseError(
            f"The PPTX has {len(slides)} slides; the limit is {options.max_pages}."
        )
    return [
        PageText(page_number=number, text=clean_text("\n".join(_shape_lines(slide.shapes))))
        for number, slide in enumerate(slides, start=1)
    ]


def _shape_lines(shapes: Iterable[BaseShape]) -> Iterator[str]:
    """Yield the text of the shapes in reading order: top to bottom, then left to right."""
    for shape in sorted(shapes, key=lambda item: (item.top or 0, item.left or 0)):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP and isinstance(shape, GroupShape):
            yield from _shape_lines(shape.shapes)
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                yield text
        elif shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    yield " | ".join(cells)